"""Document text extraction and chunking.

Supported formats:
  - PDF (text layer, with OCR fallback at 400 DPI for scanned/image-only files)
  - Word (.docx) — paragraphs + tables
  - Plain text (.txt) and Markdown (.md)
  - CSV / TSV (rows joined into "cell | cell" lines)
  - Excel (.xlsx) — every sheet, rows joined into "cell | cell" lines
"""
import io

from pypdf import PdfReader

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt", "md", "csv", "tsv", "xlsx")


# ---------------------------------------------------------------- PDF

def _text_layer_pages(file):
    """Extract per-page text from the PDF's embedded text layer (fast path)."""
    try:
        file.seek(0)  # make sure we read from the start of the stream
    except (AttributeError, OSError):
        pass
    reader = PdfReader(file)
    return [(page.extract_text() or "") for page in reader.pages]


def _ocr_pages(file, dpi=400):
    """Render each page to a high-resolution image and read it with RapidOCR.

    Used only when the text layer is empty/missing (scanned or image-based PDFs).
    dpi=400 is important: small embedded images (e.g. ticket screenshots) only
    become machine-readable when the page is rasterized at high resolution.
    Returns per-page OCR text.
    """
    try:
        import pymupdf as mupdf
    except ImportError:
        import fitz as mupdf
    from rapidocr_onnxruntime import RapidOCR
    import numpy as np

    ocr = RapidOCR()
    try:
        file.seek(0)  # the text-layer pass already consumed the stream
    except (AttributeError, OSError):
        pass
    data = file.read()
    doc = mupdf.open(stream=data, filetype="pdf")
    pages = []
    for page in doc:
        pix = page.get_pixmap(dpi=dpi)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        if pix.n == 4:
            img = img[:, :, :3]
        result, _ = ocr(img)
        pages.append("\n".join(line[1] for line in result) if result else "")
    doc.close()
    return pages


def chunk_pdf(file, chunk_size=400, overlap=50):
    """Extract text from a PDF file object and split into overlapping word chunks.

    Tries the PDF's text layer first; if it yields (almost) no text, falls back
    to OCR so scanned/image-based PDFs still work.

    Args:
        file: file-like object with .read() (e.g. Streamlit UploadedFile).
        chunk_size: target number of words per chunk.
        overlap: number of overlapping words between consecutive chunks.

    Returns:
        (chunks, ocr_used): list of non-empty text chunks, and whether OCR was needed.
    """
    pages = _text_layer_pages(file)
    ocr_used = False
    if sum(len(p.strip()) for p in pages) < 20:  # effectively no text layer
        try:
            pages = _ocr_pages(file)
            ocr_used = True
        except Exception:
            # OCR unavailable/failed — behave as before (zero chunks -> app shows guidance)
            pages = []

    text = "\n".join(pages)
    return _chunk_words(text, chunk_size, overlap), ocr_used


# ------------------------------------------------- Word / text / CSV / Excel

def _extract_docx(file):
    """Word document: paragraphs plus table rows (cells joined with ' | ')."""
    from docx import Document

    doc = Document(io.BytesIO(file.read()))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_text(file):
    """Plain text / Markdown, tolerating UTF-8 BOM and legacy encodings."""
    data = file.read()
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_delimited(file, delimiter=","):
    """CSV/TSV: every non-empty row becomes a 'cell | cell' line."""
    import csv

    text = _extract_text(file)
    lines = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [c.strip() for c in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_xlsx(file):
    """Excel workbook: every sheet, non-empty rows as 'cell | cell' lines."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        sheet_lines = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                sheet_lines.append(" | ".join(cells))
        if sheet_lines:
            parts.append(f"[Sheet: {ws.title}]")
            parts.extend(sheet_lines)
    wb.close()
    return "\n".join(parts)


# ------------------------------------------------------------- dispatcher

def _chunk_words(text, chunk_size=400, overlap=50):
    """Split text into overlapping word chunks (shared by every format)."""
    words = text.split()
    chunks = []
    step = max(1, chunk_size - overlap)
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        if i + chunk_size >= len(words):
            break
    return chunks


def chunk_file(file, filename=None, chunk_size=400, overlap=50):
    """Extract text from any supported document and split into word chunks.

    Args:
        file: file-like object with .read() (e.g. Streamlit UploadedFile).
        filename: original filename (used to pick the extractor); defaults to
            file.name when available.
        chunk_size: target number of words per chunk.
        overlap: number of overlapping words between consecutive chunks.

    Returns:
        (chunks, ocr_used): list of non-empty text chunks, and whether OCR
        was needed (always False for non-PDF formats).

    Raises:
        ValueError: if the file extension is not supported.
    """
    name = (filename or getattr(file, "name", "") or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""

    if ext == "pdf":
        return chunk_pdf(file, chunk_size=chunk_size, overlap=overlap)
    if ext == "docx":
        text = _extract_docx(file)
    elif ext in ("txt", "md", "markdown"):
        text = _extract_text(file)
    elif ext == "csv":
        text = _extract_delimited(file, delimiter=",")
    elif ext == "tsv":
        text = _extract_delimited(file, delimiter="\t")
    elif ext == "xlsx":
        text = _extract_xlsx(file)
    else:
        raise ValueError(
            f"Unsupported file type: .{ext or 'unknown'} — supported formats: "
            + ", ".join(SUPPORTED_EXTENSIONS)
        )
    return _chunk_words(text, chunk_size, overlap), False
